from __future__ import annotations

import argparse
import json
import re
import shutil
import subprocess
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from urllib.parse import urlparse
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook


API_BASE = "http://sys.xhlline.com:18083/xhl-admin"
PARTY_FIELDS = ["发货人", "收货人", "通知人"]
ROLE_LABELS = {
    "shipperRole": "发货人",
    "consigneeRole": "收货人",
    "notifyRole": "通知人",
    "forwarderRole": "货代",
    "clientRole": "客户",
}
WEAK_TOKENS = {
    "LLC",
    "LTD",
    "CO",
    "CORP",
    "COMPANY",
    "LIMITED",
    "THE",
    "AND",
    "OOO",
    "PTE",
    "INC",
    "LIABILITY",
    "JSC",
    "IP",
}
NOISY_TOKENS = {
    "CHINA",
    "NINGBO",
    "QINGDAO",
    "CIXI",
    "RUSSIA",
    "INTERNATIONAL",
    "TRADE",
}
QUERY_HEADERS = [
    "查询状态",
    "查询说明",
    "发货人匹配名称",
    "发货人匹配编码",
    "发货人匹配角色",
    "发货人查询词",
    "收货人匹配名称",
    "收货人匹配编码",
    "收货人匹配角色",
    "收货人查询词",
    "通知人匹配名称",
    "通知人匹配编码",
    "通知人匹配角色",
    "通知人查询词",
]


@dataclass
class MatchResult:
    status: str
    reason: str
    matched_names: list[str]
    matched_codes: list[str]
    matched_roles: list[str]
    attempted_terms: list[str]


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u00a0", " ").replace("\u3000", " ")
    return re.sub(r"\s+", " ", value).strip()


def split_key_value(line: str) -> tuple[str, str]:
    normalized = normalize_text(line).replace("：", ":")
    if ":" not in normalized:
        return "", ""
    key, value = normalized.split(":", 1)
    return key.strip(), value.strip()


def read_credentials(path: Path, account_index: str = "") -> tuple[str, str]:
    values: dict[str, str] = {}
    suffix = normalize_text(account_index)
    username = ""
    password = ""
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        key, value = split_key_value(raw_line)
        if key:
            values[key] = value

    if suffix:
        username = values.get(f"账号{suffix}", "")
        password = values.get(f"密码{suffix}", "")
    if not username or not password:
        username = values.get("账号", "")
        password = values.get("密码", "")
    if not username or not password:
        hint = f"账号{suffix}/密码{suffix}" if suffix else "账号/密码"
        raise RuntimeError(f"未能从 {path} 读取 {hint}")
    return username, password


def curl_post_json(
    url: str,
    token: str,
    payload: dict[str, Any],
    timeout: int,
    resolve: str = "",
) -> dict[str, Any]:
    curl = shutil.which("curl")
    if not curl:
        raise RuntimeError("curl is not available for network fallback")
    command = [
        curl,
        "-sS",
        "-m",
        str(timeout),
        "-H",
        "Content-Type: application/json",
    ]
    if resolve:
        parsed = urlparse(url)
        if parsed.hostname and parsed.port:
            command.extend(["--resolve", f"{parsed.hostname}:{parsed.port}:{resolve}"])
    if token:
        command.extend(["-H", f"Authorization: Bearer {token}"])
    command.extend(["-d", json.dumps(payload, ensure_ascii=False), url])
    completed = subprocess.run(command, check=True, capture_output=True, text=True)
    return json.loads(completed.stdout)


def post_json(
    url: str,
    token: str,
    payload: dict[str, Any],
    timeout: int = 15,
    use_curl_fallback: bool = False,
    resolve: str = "",
) -> dict[str, Any]:
    headers = {"Content-Type": "application/json"}
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except HTTPError as exc:
        body = exc.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"HTTP {exc.code} {url}: {body}") from exc
    except URLError as exc:
        if use_curl_fallback:
            try:
                return curl_post_json(url, token, payload, timeout, resolve)
            except Exception as curl_exc:
                raise RuntimeError(f"Network error for {url}: {exc}; curl fallback failed: {curl_exc}") from curl_exc
        raise RuntimeError(f"Network error for {url}: {exc}") from exc
    except TimeoutError as exc:
        if use_curl_fallback:
            try:
                return curl_post_json(url, token, payload, timeout, resolve)
            except Exception as curl_exc:
                raise RuntimeError(f"Timeout for {url}; curl fallback failed: {curl_exc}") from curl_exc
        raise RuntimeError(f"Timeout for {url}") from exc


def login(account_path: Path, account_index: str = "", use_curl_fallback: bool = False, resolve: str = "") -> str:
    username, password = read_credentials(account_path, account_index)
    result = post_json(
        f"{API_BASE}/login",
        "",
        {
            "username": username,
            "password": password,
        },
        use_curl_fallback=use_curl_fallback,
        resolve=resolve,
    )
    token = normalize_text(str(result.get("token", "")))
    if result.get("code") != 200 or not token:
        raise RuntimeError(f"登录失败: {json.dumps(result, ensure_ascii=False)}")
    return token


def strip_code_suffix(value: str) -> str:
    return re.sub(r"\s*\(EB\d+\)\s*$", "", value, flags=re.I).strip()


def normalized_compare_key(value: str) -> str:
    value = strip_code_suffix(normalize_text(value)).upper()
    return re.sub(r"[^A-Z0-9]+", "", value)


def tokenize_company(value: str) -> list[str]:
    value = strip_code_suffix(normalize_text(value)).upper()
    tokens = re.findall(r"[A-Z0-9]+", value)
    return [token for token in tokens if token and token not in WEAK_TOKENS]


def query_terms(company: str) -> list[str]:
    company = normalize_text(company)
    variants = [
        company,
        company.replace("<<", '"').replace(">>", '"'),
        company.replace("<<", "").replace(">>", ""),
    ]
    stripped = re.sub(r"[^A-Za-z0-9]+", " ", company)
    variants.append(normalize_text(stripped))
    terms: list[str] = []
    for term in variants + tokenize_company(company):
        term = normalize_text(term)
        if not term:
            continue
        if term.upper() in NOISY_TOKENS:
            continue
        if term not in terms:
            terms.append(term)
    return terms


def is_equivalent(target: str, candidate: str) -> bool:
    target_key = normalized_compare_key(target)
    candidate_key = normalized_compare_key(candidate)
    if target_key and target_key == candidate_key:
        return True

    target_tokens = set(tokenize_company(target))
    candidate_tokens = set(tokenize_company(candidate))
    if target_tokens and target_tokens.issubset(candidate_tokens):
        return True
    return False


def score_candidate(target: str, candidate: str) -> int:
    if is_equivalent(target, candidate):
        return 100
    target_tokens = set(tokenize_company(target))
    candidate_tokens = set(tokenize_company(candidate))
    if not target_tokens:
        return 0
    return len(target_tokens & candidate_tokens) * 10


def query_customer_rows(token: str, term: str, use_curl_fallback: bool = False, resolve: str = "") -> list[dict[str, Any]]:
    payload = {
        "data": {
            "nameEn": term,
            "customerType": "EB",
            "roles": [],
        },
        "pageInfo": {
            "pageNum": 1,
            "pageSize": 100,
        },
        "sortInfo": {
            "order": "desc",
            "prop": "",
        },
    }
    result = post_json(
        f"{API_BASE}/customer/listCustomer",
        token,
        payload,
        use_curl_fallback=use_curl_fallback,
        resolve=resolve,
    )
    if result.get("code") != 200:
        raise RuntimeError(f"查询失败: {json.dumps(result, ensure_ascii=False)}")
    return result.get("data", {}).get("rows", [])


def dedupe_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen: set[str] = set()
    deduped: list[dict[str, Any]] = []
    for row in rows:
        code = normalize_text(str(row.get("code", "")))
        if not code or code in seen:
            continue
        seen.add(code)
        deduped.append(row)
    return deduped


def choose_matches(company: str, rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    exact = [row for row in rows if is_equivalent(company, str(row.get("nameEn", "")))]
    if exact:
        return dedupe_rows(exact)

    scored = sorted(
        ((score_candidate(company, str(row.get("nameEn", ""))), row) for row in rows),
        key=lambda item: item[0],
        reverse=True,
    )
    best_score = scored[0][0] if scored else 0
    if best_score >= 20:
        return dedupe_rows([row for score, row in scored if score == best_score])
    return []


def collect_roles(row: dict[str, Any]) -> list[str]:
    roles: list[str] = []
    for key, label in ROLE_LABELS.items():
        if normalize_text(str(row.get(key, ""))).upper() == "Y":
            roles.append(label)
    return roles


def find_match(token: str, company: str, use_curl_fallback: bool = False, resolve: str = "") -> MatchResult:
    company = normalize_text(company)
    attempted_terms: list[str] = []
    tried_rows: list[dict[str, Any]] = []

    try:
        terms = query_terms(company)
        full_rows = query_customer_rows(token, terms[0], use_curl_fallback, resolve)
        attempted_terms.append(company)
        tried_rows.extend(full_rows)
    except RuntimeError as exc:
        return MatchResult(
            status="查询异常",
            reason=str(exc),
            matched_names=[],
            matched_codes=[],
            matched_roles=[],
            attempted_terms=[company],
        )

    matches = choose_matches(company, full_rows)
    if matches:
        return MatchResult(
            status="命中",
            reason="全名查询命中",
            matched_names=[normalize_text(str(row.get("nameEn", ""))) for row in matches],
            matched_codes=[normalize_text(str(row.get("code", ""))) for row in matches],
            matched_roles=sorted({role for row in matches for role in collect_roles(row)}),
            attempted_terms=attempted_terms,
        )

    for term in terms[1:]:
        try:
            rows = query_customer_rows(token, term, use_curl_fallback, resolve)
        except RuntimeError as exc:
            return MatchResult(
                status="查询异常",
                reason=str(exc),
                matched_names=[],
                matched_codes=[],
                matched_roles=[],
                attempted_terms=attempted_terms + [term],
            )
        attempted_terms.append(term)
        tried_rows.extend(rows)

    matches = choose_matches(company, dedupe_rows(tried_rows))
    if matches:
        return MatchResult(
            status="命中",
            reason="关键词回退命中",
            matched_names=[normalize_text(str(row.get("nameEn", ""))) for row in matches],
            matched_codes=[normalize_text(str(row.get("code", ""))) for row in matches],
            matched_roles=sorted({role for row in matches for role in collect_roles(row)}),
            attempted_terms=attempted_terms,
        )

    if tried_rows:
        return MatchResult(
            status="未命中",
            reason="有返回结果但无等价公司",
            matched_names=[],
            matched_codes=[],
            matched_roles=[],
            attempted_terms=attempted_terms,
        )

    return MatchResult(
        status="未命中",
        reason="全名和关键词查询均无结果",
        matched_names=[],
        matched_codes=[],
        matched_roles=[],
        attempted_terms=attempted_terms,
    )


def resolve_sheet_name(workbook: Any) -> str:
    for candidate in ["收发通", "查询结果"]:
        if candidate in workbook.sheetnames:
            return candidate
    return workbook.sheetnames[0]


def ensure_headers(sheet: Any) -> dict[str, int]:
    headers = [cell.value for cell in sheet[1]]
    header_to_col = {normalize_text(str(header)): idx + 1 for idx, header in enumerate(headers) if header}
    next_col = len(headers) + 1
    for header in QUERY_HEADERS:
        key = normalize_text(header)
        if key in header_to_col:
            continue
        sheet.cell(1, next_col).value = header
        header_to_col[key] = next_col
        next_col += 1
    return header_to_col


def read_cell(sheet: Any, row_idx: int, header_to_col: dict[str, int], *names: str) -> str:
    for name in names:
        key = normalize_text(name)
        if key not in header_to_col:
            continue
        value = sheet.cell(row_idx, header_to_col[key]).value
        return "" if value is None else str(value)
    return ""


def write_cell(sheet: Any, row_idx: int, header_to_col: dict[str, int], name: str, value: str) -> None:
    sheet.cell(row_idx, header_to_col[normalize_text(name)]).value = value


def latest_matching_workbook(input_path: Path | None) -> Path:
    if input_path and input_path.is_file():
        return input_path
    search_root = input_path if input_path else Path.cwd()
    if search_root.is_file():
        return search_root
    files = sorted(search_root.glob("收发通_提取信息_*.xlsx"), key=lambda path: path.stat().st_mtime)
    if not files:
        raise FileNotFoundError("未找到收发通提取 Excel")
    return files[-1]


def update_workbook(
    token: str,
    source_path: Path,
    output_path: Path,
    use_curl_fallback: bool = False,
    resolve: str = "",
) -> dict[str, int]:
    workbook = load_workbook(source_path)
    sheet = workbook[resolve_sheet_name(workbook)]
    header_to_col = ensure_headers(sheet)

    found_rows: list[list[str]] = []
    missing_rows: list[list[str]] = []
    bill_rows = 0
    hit_rows = 0
    miss_rows = 0
    error_rows = 0
    cache: dict[str, MatchResult] = {}

    for row_idx in range(2, sheet.max_row + 1):
        bl_no = read_cell(sheet, row_idx, header_to_col, "提单号", "bl_no")
        if not normalize_text(bl_no):
            continue
        bill_rows += 1
        source_file = read_cell(sheet, row_idx, header_to_col, "文件源", "source_file")
        notes: list[str] = []
        statuses: list[str] = []
        for field in PARTY_FIELDS:
            company = normalize_text(read_cell(sheet, row_idx, header_to_col, field))
            if not company:
                continue
            result = cache.get(company)
            if result is None:
                result = find_match(token, company, use_curl_fallback, resolve)
                cache[company] = result
            statuses.append(result.status)
            write_cell(sheet, row_idx, header_to_col, f"{field}匹配名称", "; ".join(result.matched_names))
            write_cell(sheet, row_idx, header_to_col, f"{field}匹配编码", "; ".join(result.matched_codes))
            write_cell(sheet, row_idx, header_to_col, f"{field}匹配角色", "/".join(result.matched_roles))
            write_cell(sheet, row_idx, header_to_col, f"{field}查询词", " | ".join(result.attempted_terms))
            notes.append(f"{field}:{result.reason}")

            summary_row = [
                bl_no,
                source_file,
                field,
                company,
                "; ".join(result.matched_names),
                "; ".join(result.matched_codes),
                "/".join(result.matched_roles),
                " | ".join(result.attempted_terms),
                result.reason,
            ]
            if result.status == "命中":
                hit_rows += 1
                found_rows.append(summary_row)
            elif result.status == "查询异常":
                error_rows += 1
                missing_rows.append(summary_row)
            else:
                miss_rows += 1
                missing_rows.append(summary_row)

        if statuses and all(status == "命中" for status in statuses):
            status_value = "查询完成"
        elif "查询异常" in statuses:
            status_value = "存在异常"
        elif "命中" in statuses:
            status_value = "部分命中"
        else:
            status_value = "未命中"

        write_cell(sheet, row_idx, header_to_col, "查询状态", status_value)
        write_cell(sheet, row_idx, header_to_col, "查询说明", "；".join(notes))

    for sheet_name in ["可查询到", "查询不到"]:
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    found_sheet = workbook.create_sheet("可查询到")
    missing_sheet = workbook.create_sheet("查询不到")
    summary_headers = ["提单号", "文件源", "字段", "公司抬头", "匹配名称", "匹配编码", "确认角色", "查询词", "说明"]
    found_sheet.append(summary_headers)
    missing_sheet.append(summary_headers)
    for row in found_rows:
        found_sheet.append(row)
    for row in missing_rows:
        missing_sheet.append(row)

    workbook.save(output_path)
    return {
        "bill_rows": bill_rows,
        "hit_rows": hit_rows,
        "miss_rows": miss_rows,
        "error_rows": error_rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path)
    parser.add_argument("--input-dir", type=Path)
    parser.add_argument("--account", required=True, type=Path)
    parser.add_argument("--account-index", default="", help="Read 账号N/密码N, for example --account-index 2")
    parser.add_argument("--curl-fallback", action="store_true", help="Use curl when Python urllib cannot reach the site")
    parser.add_argument("--resolve-ip", default="", help="Optional IP for curl --resolve host:port:ip fallback")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    workbook_path = latest_matching_workbook(args.workbook or args.input_dir)
    output_path = args.output or workbook_path.parent / f"查询比对_结果_{datetime.now().strftime('%Y%m%d-%H%M')}.xlsx"
    token = login(args.account, args.account_index, args.curl_fallback, args.resolve_ip)
    stats = update_workbook(token, workbook_path, output_path, args.curl_fallback, args.resolve_ip)
    print(workbook_path)
    print(output_path)
    print(json.dumps(stats, ensure_ascii=False))


if __name__ == "__main__":
    main()
